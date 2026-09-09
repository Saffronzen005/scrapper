import pyautogui
import openpyxl
from playwright.sync_api import sync_playwright
from seleniumbase import sb_cdp

# --- 1. Define the Webshop Function to Return Data ---
def elevator_parts_eu(pw, code):
    print(f"\n--- Extracting {code} from Elevator-Parts.eu ---")
    
    browser = pw.chromium.launch(headless=False, slow_mo=1000)
    context = browser.new_context(viewport={"width": 1920, "height": 1080})
    page = context.new_page()
    
    page.goto("https://elevator-parts.eu/index.php?route=product/catalog&page=3055")
    page.get_by_placeholder("Search here...").fill(code)
    page.get_by_placeholder("Search here...").press("Enter")
    
    page.wait_for_load_state("domcontentloaded")

    # Check if product exists on the webshop
    no_results_message = page.get_by_text("There is no product that matches the search criteria.")
    if no_results_message.is_visible():
        browser.close()
        return {"Availability": "Not Available"}

    # Click the product and extract data
    page.locator(f".name a:has-text('{code}')").first.click()
    page.wait_for_load_state("domcontentloaded")
    
    title = page.title()
    url = page.url
    raw_price = page.locator('.product-price').inner_text() # e.g., "0.00€"
    
    stock_value = page.locator(".product-stock span").inner_text()
    brand_value = page.locator(".product-manufacturer a").inner_text()
    model_value = page.locator(".product-model span").inner_text()

    # Format the data to match your Excel columns
    price = raw_price.replace('€', '').strip()
    additional_text = f"{title} Stock:{stock_value}Brand:{brand_value}Model:{model_value}"
    
    browser.close()
    
    # Return as a dictionary for the main loop to write to Excel
    return {
        "Availability": "In stock",
        "Additional Text": additional_text,
        "Supplier": "Elevator-parts.eu",
        "WebShop url": url,
        "Price": price,
        "Currency": "EUR" if "€" in raw_price else ""
    }


# --- 2. Main Execution: Excel Processing ---
EXCEL_FILE = "book.xlsx" # Change this to your actual file name

# Load the Excel workbook and select the active sheet
wb = openpyxl.load_workbook(EXCEL_FILE)
sheet = wb.active

sb = sb_cdp.Chrome()
endpoint_url = sb.get_endpoint_url()

with sync_playwright() as pw:
    # Loop through rows in Excel (starting at row 2 to skip headers)
    for row in range(2, sheet.max_row + 1):
        # Read the material code from Column A (Column 1)
        code = sheet.cell(row=row, column=1).value
        
        if not code:
            continue # Skip empty rows
            
        print(f"\nProcessing row {row}: {code}")
        
        browser = pw.chromium.connect_over_cdp(endpoint_url)
        context = browser.contexts[0]
        page = context.pages[0]

        query = f'"{code}" -parts.kone.com -scribd.com -list -elespares.com'
        
        page.goto("https://www.google.com")
        sb.sleep(2)
        sb.solve_captcha()
        sb.wait_for_element_absent("input[disabled]")
        sb.sleep(2)

        search_box = page.locator('textarea[name="q"]')
        search_box.fill(query)
        search_box.press("Enter")
        page.wait_for_load_state("domcontentloaded")
        
        # Check Google results
        not_avai = page.get_by_text("did not match any documents.")
        
        extracted_data = {}
        if not_avai.is_visible():
            print(f"Product '{code}' Not available on Google.")
            extracted_data = {"Availability": "Not Available"}
        else:
            names_locator = page.locator('.VuuXrf')
            all_names = set(names_locator.all_inner_texts())

            if "elevator-parts.eu" in all_names:
                # Call function and get the returned dictionary
                extracted_data = elevator_parts_eu(pw, code)
            else:
                extracted_data = {"Availability": "Not Available"}
                
        browser.close()

        # --- 3. Write Data Back to Excel ---
        # Map the extracted dictionary values to the correct Excel columns
        sheet.cell(row=row, column=2).value = extracted_data.get("Availability", "")
        
        if extracted_data.get("Availability") != "Not Available":
            sheet.cell(row=row, column=3).value = extracted_data.get("Additional Text", "")
            sheet.cell(row=row, column=4).value = extracted_data.get("Supplier", "")
            sheet.cell(row=row, column=5).value = extracted_data.get("WebShop url", "")
            sheet.cell(row=row, column=6).value = extracted_data.get("Price", "")
            sheet.cell(row=row, column=7).value = extracted_data.get("Currency", "")

        # Save after every row so you don't lose data if the script crashes
        wb.save(EXCEL_FILE) 

print("\nFinished processing all rows and saved to Excel!")