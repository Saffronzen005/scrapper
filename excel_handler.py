"""
Excel I/O handler — reads product codes and writes styled output reports.

Handles:
  - Reading material codes from the input Excel
  - Creating a styled output workbook
  - Appending result rows incrementally
  - Generating a summary sheet
"""

import os
import logging
from datetime import datetime

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

logger = logging.getLogger(__name__)


# ── Styling constants ────────────────────────────────────────────────
HEADER_FONT = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
HEADER_FILL = PatternFill(start_color="2F5496", end_color="2F5496", fill_type="solid")
HEADER_ALIGNMENT = Alignment(horizontal="center", vertical="center", wrap_text=True)

UNAUTHORIZED_FILL = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
UNAUTHORIZED_FONT = Font(name="Calibri", color="9C0006")

CLEAN_FILL = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
CLEAN_FONT = Font(name="Calibri", color="006100")

ERROR_FILL = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
ERROR_FONT = Font(name="Calibri", color="9C6500")

THIN_BORDER = Border(
    left=Side(style="thin", color="D9D9D9"),
    right=Side(style="thin", color="D9D9D9"),
    top=Side(style="thin", color="D9D9D9"),
    bottom=Side(style="thin", color="D9D9D9"),
)

# Output column headers
OUTPUT_HEADERS = [
    "Material Code",
    "Status",
    "Seller Domain",
    "Seller URL",
    "Page Title",
    "Price",
    "Currency",
    "Stock Status",
    "Search Snippet",
    "Scrape Error",
]

COLUMN_WIDTHS = [18, 22, 30, 55, 45, 12, 10, 16, 50, 25]


def read_product_codes(filepath: str) -> list[str]:
    """
    Read product/material codes from column A of the input Excel file.

    Args:
        filepath: Path to the input .xlsx file

    Returns:
        List of product code strings (skipping the header row and empty cells)
    """
    wb = openpyxl.load_workbook(filepath, read_only=True)
    sheet = wb.active
    codes = []

    for row in sheet.iter_rows(min_row=2, max_col=1, values_only=True):
        value = row[0]
        if value is not None:
            codes.append(str(value).strip())

    wb.close()
    logger.info(f"Loaded {len(codes)} product codes from '{filepath}'")
    return codes


def create_output_workbook(output_path: str) -> openpyxl.Workbook:
    """
    Create a new styled output workbook with headers.

    Args:
        output_path: Where to save the workbook

    Returns:
        The openpyxl Workbook object (caller is responsible for saving)
    """
    wb = openpyxl.Workbook()
    sheet = wb.active
    sheet.title = "Unauthorized Sellers"

    # Write and style headers
    for col_idx, header in enumerate(OUTPUT_HEADERS, 1):
        cell = sheet.cell(row=1, column=col_idx, value=header)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = HEADER_ALIGNMENT
        cell.border = THIN_BORDER

    # Set column widths
    for col_idx, width in enumerate(COLUMN_WIDTHS, 1):
        sheet.column_dimensions[get_column_letter(col_idx)].width = width

    # Freeze the header row
    sheet.freeze_panes = "A2"

    # Auto-filter
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(OUTPUT_HEADERS))}1"

    # Ensure output directory exists
    os.makedirs(os.path.dirname(output_path) or ".", exist_ok=True)
    wb.save(output_path)

    logger.info(f"Created output workbook: '{output_path}'")
    return wb


def load_output_workbook(output_path: str) -> openpyxl.Workbook:
    """Load an existing output workbook (for resume)."""
    wb = openpyxl.load_workbook(output_path)
    logger.info(f"Resumed output workbook: '{output_path}' ({wb.active.max_row - 1} existing rows)")
    return wb


def append_unauthorized_row(
    sheet,
    product_code: str,
    domain: str,
    url: str,
    title: str,
    price: str,
    currency: str,
    stock_status: str,
    snippet: str,
    scrape_error: str = "",
):
    """Append a row for an unauthorized seller finding."""
    row_num = sheet.max_row + 1
    values = [
        product_code,
        "⚠️ UNAUTHORIZED",
        domain,
        url,
        title,
        price,
        currency,
        stock_status,
        snippet[:200] if snippet else "",
        scrape_error,
    ]

    for col_idx, value in enumerate(values, 1):
        cell = sheet.cell(row=row_num, column=col_idx, value=value)
        cell.font = UNAUTHORIZED_FONT
        cell.fill = UNAUTHORIZED_FILL
        cell.border = THIN_BORDER
        if col_idx in (4, 5, 9):  # URL, title, snippet columns — wrap text
            cell.alignment = Alignment(wrap_text=True, vertical="top")


def append_clean_row(sheet, product_code: str):
    """Append a row for a product with no unauthorized sellers found."""
    row_num = sheet.max_row + 1
    values = [
        product_code,
        "✅ CLEAN",
        "", "", "", "", "", "", "", "",
    ]

    for col_idx, value in enumerate(values, 1):
        cell = sheet.cell(row=row_num, column=col_idx, value=value)
        cell.font = CLEAN_FONT
        cell.fill = CLEAN_FILL
        cell.border = THIN_BORDER


def append_error_row(sheet, product_code: str, error_msg: str):
    """Append a row for a product that encountered a processing error."""
    row_num = sheet.max_row + 1
    values = [
        product_code,
        "❌ ERROR",
        "", "", "", "", "", "", "",
        error_msg[:200],
    ]

    for col_idx, value in enumerate(values, 1):
        cell = sheet.cell(row=row_num, column=col_idx, value=value)
        cell.font = ERROR_FONT
        cell.fill = ERROR_FILL
        cell.border = THIN_BORDER


def append_no_results_row(sheet, product_code: str):
    """Append a row for a product with no search results at all."""
    row_num = sheet.max_row + 1
    values = [
        product_code,
        "🔍 NO RESULTS",
        "", "", "", "", "", "", "", "",
    ]

    for col_idx, value in enumerate(values, 1):
        cell = sheet.cell(row=row_num, column=col_idx, value=value)
        cell.border = THIN_BORDER


def add_summary_sheet(wb: openpyxl.Workbook, stats: dict):
    """
    Add a summary sheet to the workbook with scan statistics.

    Args:
        wb:    The output workbook
        stats: Dictionary with scan statistics
    """
    if "Summary" in wb.sheetnames:
        del wb["Summary"]

    ws = wb.create_sheet("Summary", 0)  # Insert at the beginning

    # Title
    ws.merge_cells("A1:D1")
    title_cell = ws["A1"]
    title_cell.value = "KONE Unauthorized Seller Scan — Summary"
    title_cell.font = Font(name="Calibri", bold=True, size=14, color="2F5496")
    title_cell.alignment = Alignment(horizontal="center")

    # Timestamp
    ws["A2"] = f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"
    ws["A2"].font = Font(name="Calibri", italic=True, color="808080")

    # Stats table
    stat_rows = [
        ("Total Products Scanned", stats.get("total", 0)),
        ("Products with Unauthorized Sellers", stats.get("unauthorized", 0)),
        ("Clean Products", stats.get("clean", 0)),
        ("No Search Results", stats.get("no_results", 0)),
        ("Errors", stats.get("errors", 0)),
        ("", ""),
        ("Unique Unauthorized Domains Found", stats.get("unique_domains", 0)),
    ]

    start_row = 4
    ws.column_dimensions["A"].width = 40
    ws.column_dimensions["B"].width = 15

    for i, (label, value) in enumerate(stat_rows):
        row = start_row + i
        ws.cell(row=row, column=1, value=label).font = Font(name="Calibri", bold=True)
        ws.cell(row=row, column=2, value=value).font = Font(name="Calibri")
        ws.cell(row=row, column=2).alignment = Alignment(horizontal="right")

    # Top unauthorized domains
    top_domains = stats.get("top_domains", [])
    if top_domains:
        domain_start = start_row + len(stat_rows) + 1
        ws.cell(row=domain_start, column=1, value="Top Unauthorized Seller Domains").font = Font(
            name="Calibri", bold=True, size=12, color="9C0006"
        )

        ws.cell(row=domain_start + 1, column=1, value="Domain").font = HEADER_FONT
        ws.cell(row=domain_start + 1, column=1).fill = HEADER_FILL
        ws.cell(row=domain_start + 1, column=2, value="Products Listed").font = HEADER_FONT
        ws.cell(row=domain_start + 1, column=2).fill = HEADER_FILL

        for i, (domain, count) in enumerate(top_domains):
            row = domain_start + 2 + i
            ws.cell(row=row, column=1, value=domain).font = UNAUTHORIZED_FONT
            ws.cell(row=row, column=1).fill = UNAUTHORIZED_FILL
            ws.cell(row=row, column=2, value=count).font = UNAUTHORIZED_FONT
            ws.cell(row=row, column=2).fill = UNAUTHORIZED_FILL
            ws.cell(row=row, column=2).alignment = Alignment(horizontal="right")
